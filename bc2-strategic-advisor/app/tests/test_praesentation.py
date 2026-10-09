"""
Tests für den Foliengenerator (#257, Format und Weg entschieden in #244).

**Geprüft wird, was auf den Folien steht, nicht dass die Datei entsteht.** Eine
PPTX ist ein ZIP mit Zeitstempeln; ein Bytevergleich taugt nicht. Drei
Prüfungen aus #244, dazu die Darstellungsregeln aus #257:

(a) **kein ``{{``** im ganzen Foliensatz — fängt das Template-Erbe ab;
(b) **Vollzähligkeit** — je freigegebenem Potenzial eine Detailfolie, je
    Kernprozess sein Abschnitt, nichts stillschweigend abgeschnitten;
(c) **die Scheingenauigkeits-Invariante** über alle Textrahmen: kein
    Euro-Betrag ohne Spanne, keine 0 als Ersatz für eine fehlende Größe.

Gelesen wird über alle Textrahmen **und** die Notizen. Die Notizen tragen je
Detail- und Kernprozessfolie die Kennung (``potenzial_id: …``, ``kp_id: …``):
daran zählt (b), statt an Titeln, die sich wiederholen dürfen.

Eingänge: die v3.1-Fixtures aus ``contracts/examples`` (volle Texte, zwei
Kernprozesse) und der Messsatz aus #167 über die Oberflächenquelle (elf
Potenziale in vier Kernprozessen, ohne LLM-Texte — der Weg, den der Knopf heute
geht).
"""

from __future__ import annotations

import copy
import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx")

from laeufe import MesssatzLaufquelle, aus_lauf  # noqa: E402
from modell import Nutzwert, Nutzwertkategorie, Potenzialeingang, rechne_lauf  # noqa: E402
from praesentation import (  # noqa: E402
    DATEINAME,
    KeineFreigabe,
    als_bytes,
    baue_praesentation,
    lege_ab,
    lieferordner,
    mandantenkuerzel,
)
from praesentation.folien import KACHELN_JE_FOLIE, ZEILEN_KOSTEN  # noqa: E402

WURZEL = Path(__file__).resolve().parents[3]
BEISPIELE = WURZEL / "contracts" / "examples"
MESSSAETZE = WURZEL / "bc2-strategic-advisor" / "kalibrierung"


# ---------------------------------------------------------------------------
# Eingänge
# ---------------------------------------------------------------------------


def _lies(name: str) -> dict:
    return json.loads((BEISPIELE / name).read_text(encoding="utf-8"))


def _freigabe(prio: dict, **gate1) -> dict:
    prio = copy.deepcopy(prio)
    prio["gate1"] = {
        "status": "approved",
        "entscheider": "Sergio (Test)",
        "entschieden_am": "2026-10-09T10:00:00+00:00",
        "approved_potenzial_ids": [e["potenzial_id"] for e in prio["eintraege"]],
        **gate1,
    }
    return prio


@pytest.fixture
def fixture_lauf() -> tuple[list[dict], dict]:
    """Die v3.1-Fixtures: zwei Konzepte (KP-05, KP-06), sechs Potenziale, freigegeben."""
    konzepte = [_lies("mock_automatisierungskonzept.json"),
                _lies("mock_automatisierungskonzept_KP-05.json")]
    return konzepte, _freigabe(_lies("mock_prozesspriorisierung.json"))


@pytest.fixture
def messsatz_lauf() -> tuple[list[dict], dict, str | None]:
    """Der Messsatz aus #167 über die Oberflächenquelle — elf Potenziale, vier Kernprozesse."""
    quelle = MesssatzLaufquelle(MESSSAETZE, uebergeben_am=datetime(2026, 9, 21, tzinfo=timezone.utc))
    ansicht = quelle.ansicht(quelle.uebersicht()[0].paket_id)
    gate1 = {
        "status": "approved",
        "entschieden_am": "2026-10-09T10:00:00+00:00",
        "approved_potenzial_ids": ansicht.potenzial_ids(),
    }
    konzepte, prio = ansicht.als_vertrag(gate1)
    return konzepte, prio, ansicht.kopf.warnung


# ---------------------------------------------------------------------------
# Lesen
# ---------------------------------------------------------------------------


def _texte_der_folie(folie) -> list[str]:
    return [sp.text_frame.text for sp in folie.shapes if sp.has_text_frame and sp.text_frame.text]


def _notiz(folie) -> str:
    return folie.notes_slide.notes_text_frame.text if folie.has_notes_slide else ""


def _alle_texte(prs) -> list[str]:
    texte = []
    for folie in prs.slides:
        texte.extend(_texte_der_folie(folie))
        if _notiz(folie):
            texte.append(_notiz(folie))
    return texte


def _volltext(prs) -> str:
    return "\n".join(_alle_texte(prs))


def _folien_mit_notiz(prs, praefix: str) -> list[str]:
    werte = []
    for folie in prs.slides:
        for zeile in _notiz(folie).splitlines():
            if zeile.startswith(praefix):
                werte.append(zeile.removeprefix(praefix).strip())
    return werte


def _kopfzeilen(prs) -> list[str]:
    """Alle Textrahmen je Folie, zu einem Text gefügt — darin stehen die Folientitel."""
    return ["\n".join(_texte_der_folie(f)) for f in prs.slides]


# ---------------------------------------------------------------------------
# Die Invariante (c), einmal formuliert und überall angewandt
# ---------------------------------------------------------------------------

_EURO = re.compile(r"(\d[\d.]*(?:,\d+)?)\s*€")
_SPANNE_DAVOR = re.compile(r"\d[\d.]*(?:,\d+)?–$")
_RICHTWERT_DAVOR = re.compile(r"Richtwert $")
_RICHTWERT_DANACH = re.compile(r"\s*aus \d[\d.]*(?:,\d+)? PT Aufwand")
_ERSATZ_NULL = re.compile(r"(?<![\d.,–])0(?:,0+)?\s*(?:€|h/Jahr|h\b|PT\b|Monate)")


def _verstoesse(prs) -> list[str]:
    """Jeder Euro-Betrag ohne Spanne und jede 0 als Ersatz — als Klartext."""
    funde = []
    for text in _alle_texte(prs):
        for m in _EURO.finditer(text):
            davor, danach = text[: m.start(1)], text[m.end():]
            if _SPANNE_DAVOR.search(davor):
                continue  # oberes Ende von „a–b €"
            if _RICHTWERT_DAVOR.search(davor) and _RICHTWERT_DANACH.match(danach):
                continue  # „Richtwert 12.000 € aus 15 PT Aufwand" — die eine benannte Form
            funde.append(f"Euro ohne Spanne: …{text[max(0, m.start() - 30): m.end() + 20]}…")
        for m in _ERSATZ_NULL.finditer(text):
            funde.append(f"0 als Ersatz: …{text[max(0, m.start() - 30): m.end() + 10]}…")
        if text.strip() in {"-", "–", "—", "0", "n/a", "N/A"} or "n/a" in text.lower():
            funde.append(f"Platzhalter statt Grund: {text!r}")
    return funde


def test_die_invariante_selbst_faengt_was_sie_fangen_soll():
    """Ohne diese Gegenprobe wäre ein grüner Test (c) auch ein blinder."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    folie = prs.slides.add_slide(prs.slide_layouts[6])
    for t in ("Einsparung 23.500 €/Jahr", "Aufwand heute: 0 h/Jahr", "n/a",
              "18.000–29.000 €/Jahr (Dauer geschätzt, ±40 %)",
              "Investition: Richtwert 12.000 € aus 15 PT Aufwand"):
        folie.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text_frame.text = t
    funde = _verstoesse(prs)
    assert len(funde) == 3, funde


# ---------------------------------------------------------------------------
# (a) Kein Template-Erbe
# ---------------------------------------------------------------------------


def test_a_kein_platzhalter_im_foliensatz(fixture_lauf, messsatz_lauf):
    for konzepte, prio, *_ in (fixture_lauf, messsatz_lauf):
        assert "{{" not in _volltext(baue_praesentation(konzepte, prio))


# ---------------------------------------------------------------------------
# (b) Vollzähligkeit
# ---------------------------------------------------------------------------


def test_b_je_potenzial_eine_detailfolie_je_kernprozess_ein_abschnitt(fixture_lauf):
    konzepte, prio = fixture_lauf
    prs = baue_praesentation(konzepte, prio)

    detail = _folien_mit_notiz(prs, "potenzial_id:")
    assert sorted(detail) == sorted(e["potenzial_id"] for e in prio["eintraege"])
    assert len(detail) == len(set(detail))

    abschnitte = _folien_mit_notiz(prs, "kp_id:")
    # Jede Detailfolie nennt ihren Kernprozess, dazu je Kernprozess die Abschnittsfolie.
    assert sorted(set(abschnitte)) == sorted(r["kp_id"] for r in prio["prozess_raenge"])


def test_b_ueberlauf_wird_umgebrochen_nie_begrenzt(messsatz_lauf):
    """Elf Potenziale gegen sechs Kacheln: zwei Überblicksfolien, nicht „Top 6"."""
    konzepte, prio, _ = messsatz_lauf
    n = len(prio["eintraege"])
    assert n > KACHELN_JE_FOLIE and n > ZEILEN_KOSTEN, "Der Messsatz prüft den Überlauf nicht mehr."
    prs = baue_praesentation(konzepte, prio)

    assert len(_folien_mit_notiz(prs, "potenzial_id:")) == n
    kopf = "\n".join(_kopfzeilen(prs))
    seiten = -(-n // KACHELN_JE_FOLIE)
    for i in range(1, seiten + 1):
        assert f"Überblick: freigegebene Potenziale ({i} von {seiten})" in kopf
    kosten = -(-n // ZEILEN_KOSTEN)
    assert f"Kosten und Nutzen je Potenzial ({kosten} von {kosten})" in kopf

    text = _volltext(prs)
    for e in prio["eintraege"]:
        # Jeder Titel im Überblick, in der Matrix, auf der Detailfolie und in der Kostentabelle.
        assert text.count(e["titel"]) >= 4, e["titel"]


def test_b_jeder_kernprozess_des_messsatzes_hat_seinen_abschnitt(messsatz_lauf):
    konzepte, prio, _ = messsatz_lauf
    prs = baue_praesentation(konzepte, prio)
    erwartet = [r["kp_id"] for r in sorted(prio["prozess_raenge"], key=lambda r: r["prozessrang"])]
    abschnitte = []
    for folie in prs.slides:
        zeilen = _notiz(folie).splitlines()
        if len(zeilen) == 1 and zeilen[0].startswith("kp_id:"):
            abschnitte.append(zeilen[0].removeprefix("kp_id:").strip())
    assert abschnitte == erwartet


# ---------------------------------------------------------------------------
# (c) Scheingenauigkeit
# ---------------------------------------------------------------------------


def test_c_kein_euro_ohne_spanne_und_keine_0_als_ersatz(fixture_lauf, messsatz_lauf):
    for konzepte, prio, *_ in (fixture_lauf, messsatz_lauf):
        assert _verstoesse(baue_praesentation(konzepte, prio)) == []


def test_c_ohne_wertaussage_steht_der_grund(keine_wertaussage_lauf):
    konzepte, prio = keine_wertaussage_lauf
    prs = baue_praesentation(konzepte, prio)
    text = _volltext(prs)
    assert _verstoesse(prs) == []
    assert "Keine Wertaussage" in text or "Warum keine Wertaussage" in text
    # Der Vertrag trägt hier stunden_jahr: 0 — auf der Folie steht der Grund.
    assert "keine Dauer erhoben" in text


def test_die_herkunft_steht_im_selben_satz_wie_die_spanne(fixture_lauf):
    konzepte, prio = fixture_lauf
    text = _volltext(baue_praesentation(konzepte, prio))
    spannen = re.findall(r"\d[\d.]*–\d[\d.]* €/Jahr \(([^)]*)\)", text)
    assert spannen, "Keine Einsparungsspanne gefunden."
    assert all("Dauer" in h and "±" in h for h in spannen), spannen


# ---------------------------------------------------------------------------
# Die übrigen Darstellungsregeln aus #257
# ---------------------------------------------------------------------------


def test_der_prioritaetsscore_erscheint_nicht(fixture_lauf):
    """Eine 74,5 sieht gemessen aus und ist ein Produkt zweier Urteile (#244)."""
    konzepte, prio = fixture_lauf
    assert "score" not in _volltext(baue_praesentation(konzepte, prio)).lower()


def test_die_titelfolie_traegt_mandant_paket_fassung_und_uebergabe(fixture_lauf):
    konzepte, prio = fixture_lauf
    titel = "\n".join(_texte_der_folie(baue_praesentation(konzepte, prio).slides[0]))
    assert "NoroAI Consulting GmbH" in titel
    assert prio["paket_id"] in titel
    assert f"Fassung {prio['fassung']}" in titel
    assert "übergeben am 31.08.2026" in titel


def test_die_ausgangslage_kommt_aus_dem_vertrag(fixture_lauf):
    konzepte, prio = fixture_lauf
    text = _volltext(baue_praesentation(konzepte, prio))
    u = prio["ausgangslage"]["unternehmen"]
    for wert in (u["branche"], u["region"], u["geschaeftsmodell"], f"Mitarbeitende: {u['mitarbeitende']}"):
        assert str(wert) in text
    for h in prio["ausgangslage"]["herausforderungen"]:
        assert h["beschreibung"] in text
    # Die Kausalketten-Folie ist bewusst entfallen (#244).
    assert "Kausalkette" not in text


def test_ohne_ausgangslage_steht_der_grund_nicht_eine_erfindung(messsatz_lauf):
    konzepte, prio, _ = messsatz_lauf
    assert "ausgangslage" not in prio and prio["schema_version"] == "3.0"
    text = _volltext(baue_praesentation(konzepte, prio))
    assert text.count("in dieser Priorisierung nicht vor") >= 2  # Unternehmen und Herausforderungen


def test_eine_kernaussage_erscheint_wenn_es_sie_gibt(fixture_lauf):
    konzepte, prio = fixture_lauf
    prio["ausgangslage"]["kernaussage"] = "Der größte Hebel liegt im Medienbruch zur Excel-Liste."
    assert "Der größte Hebel liegt im Medienbruch" in _volltext(baue_praesentation(konzepte, prio))


def test_querschnitte_stehen_auf_der_detailfolie(fixture_lauf):
    konzepte, prio = fixture_lauf
    prs = baue_praesentation(konzepte, prio)
    pot = konzepte[0]["potenziale"][0]
    for folie in prs.slides:
        if f"potenzial_id: {pot['potenzial_id']}" in _notiz(folie):
            assert pot["querschnitte"]["zukunftssicherheit"] in "\n".join(_texte_der_folie(folie))
            return
    pytest.fail("Detailfolie nicht gefunden.")


def test_nicht_freigegebene_kommen_nicht_vor_auch_nicht_ihr_grund(fixture_lauf):
    konzepte, prio = fixture_lauf
    raus = prio["eintraege"][-1]
    grund = "Passt nicht in das Budget dieses Jahres — erst nach dem Umzug."
    prio["gate1"]["approved_potenzial_ids"].remove(raus["potenzial_id"])
    prio["gate1"]["nicht_freigegeben"] = [{"potenzial_id": raus["potenzial_id"], "begruendung": grund}]

    prs = baue_praesentation(konzepte, prio)
    text = _volltext(prs)
    assert raus["titel"] not in text
    assert grund not in text
    assert raus["potenzial_id"] not in _folien_mit_notiz(prs, "potenzial_id:")


def test_die_am_gate_gesetzte_reihenfolge_gilt(fixture_lauf):
    konzepte, prio = fixture_lauf
    folge = [e["potenzial_id"] for e in reversed(prio["eintraege"])]
    prio["gate1"]["finale_reihenfolge_potenzial_ids"] = folge
    prio["gate1"]["abweichungsbegruendung"] = "Das Budget erlaubt nur die kleinen zuerst."
    prs = baue_praesentation(konzepte, prio)
    text = _volltext(prs)
    letzter = prio["eintraege"][-1]
    assert f"Rang 1   {letzter['titel']}" in text
    assert "Das Budget erlaubt nur die kleinen zuerst." in text


@pytest.mark.parametrize("status", ["pending", "rejected"])
def test_vor_der_freigabe_und_bei_ablehnung_entsteht_keine_praesentation(fixture_lauf, status):
    konzepte, prio = fixture_lauf
    prio["gate1"] = {"status": status, "kommentar": "nicht freigegeben"}
    with pytest.raises(KeineFreigabe):
        baue_praesentation(konzepte, prio)


def test_die_warnung_der_quelle_steht_auf_der_titelfolie(messsatz_lauf):
    konzepte, prio, warnung = messsatz_lauf
    assert warnung, "Der Messsatz trägt keine Warnung mehr — dann prüft dieser Test nichts."
    titel = "\n".join(_texte_der_folie(baue_praesentation(konzepte, prio, warnung=warnung).slides[0]))
    assert warnung in titel


def test_derselbe_lauf_ergibt_dieselben_folien(fixture_lauf):
    """Reproduzierbar im Inhalt — der Bytevergleich taugt nicht (Zeitstempel im ZIP)."""
    konzepte, prio = fixture_lauf
    assert _alle_texte(baue_praesentation(konzepte, prio)) == _alle_texte(
        baue_praesentation(konzepte, prio))


# ---------------------------------------------------------------------------
# Alle ohne Wertaussage — die offene Frage aus #257
# ---------------------------------------------------------------------------


def _nutzwert() -> Nutzwert:
    k = Nutzwertkategorie(6, "Begruendungssatz fuer den Test.")
    return Nutzwert(k, k, k, k, k)


@pytest.fixture
def keine_wertaussage_lauf() -> tuple[list[dict], dict]:
    """Drei Potenziale, keines mit erhobener Dauer — ``value_quelle: "keine"`` überall."""
    eingaenge = [
        Potenzialeingang(
            potenzial_id=f"00000000-0000-4000-8000-00000000000{i}",
            titel=f"Potenzial ohne Dauer {i}",
            kp_id="KP-02",
            betroffene_teilprozess_ids=(f"KP-02.TP-{i}",),
            klasse="Integration",
            automatisierungsgrad_begruendung="Zwei Systeme verbinden.",
            nutzwert=_nutzwert(),
            frequency_per_year=None,
            total_duration_minutes=None,
            focus_step_duration_source=None,
            reifeskalen=(3, 3, 3, 3),
            aufwand_schaetzung_pt=None if i == 1 else 8.0,
        )
        for i in (1, 2, 3)
    ]
    lauf = rechne_lauf("7c2d5ee9-2a9a-5990-810f-502ea2b2012d", "PKT-OHNE-DAUER", eingaenge)
    ansicht = aus_lauf(lauf, eingaenge, uebergeben_am=datetime(2026, 9, 21, tzinfo=timezone.utc))
    return ansicht.als_vertrag({"status": "approved", "approved_potenzial_ids": ansicht.potenzial_ids()})


def test_sind_alle_ohne_wertaussage_zeigt_teil_3_gruende_statt_leerer_tabellen(keine_wertaussage_lauf):
    konzepte, prio = keine_wertaussage_lauf
    assert all(p["value"]["value_quelle"] == "keine" for k in konzepte for p in k["potenziale"])
    prs = baue_praesentation(konzepte, prio)
    text = _volltext(prs)
    assert "Für keines der 3 Potenziale liegt eine Wertaussage vor" in text
    assert "Warum keine Wertaussage" in text
    assert "Einsparung je Jahr" not in text
    assert "kein Aufwand geschätzt, deshalb kein Richtwert" in text


# ---------------------------------------------------------------------------
# Ablage
# ---------------------------------------------------------------------------


def test_die_ablage_schreibt_in_den_lieferordner_ohne_fassung_im_dateinamen(tmp_path, fixture_lauf):
    konzepte, prio = fixture_lauf
    daten = als_bytes(baue_praesentation(konzepte, prio))
    basis = tmp_path / "lieferungen"
    basis.mkdir()
    ordner = lieferordner(basis, mandantenkuerzel("NoroAI Consulting GmbH", prio["company_id"]),
                          prio["paket_id"], prio["fassung"])
    pfad = lege_ab(daten, ordner)

    assert pfad == basis / f"noroai-{prio['paket_id']}-f1" / DATEINAME
    assert DATEINAME == "praesentation.pptx"
    gelesen = pptx.Presentation(io.BytesIO(pfad.read_bytes()))
    assert len(gelesen.slides) == len(baue_praesentation(konzepte, prio).slides)


def test_ohne_lieferungen_verzeichnis_wird_nichts_erfunden(tmp_path):
    with pytest.raises(FileNotFoundError):
        lege_ab(b"PK", tmp_path / "gibt-es-nicht" / "noroai-PKT-f1")


def test_das_mandantenkuerzel():
    assert mandantenkuerzel("NoroAI Consulting GmbH", "7c2d5ee9-x") == "noroai"
    assert mandantenkuerzel(None, "7C2D5EE9-2a9a") == "7c2d5ee9"
    assert mandantenkuerzel("  ", "7c2d5ee9-2a9a") == "7c2d5ee9"


# ---------------------------------------------------------------------------
# Kopf der Detailfolie: der Titel läuft nie unter das PRIO-Etikett
# ---------------------------------------------------------------------------


def _titel_und_etikett(folie, titel: str):
    rahmen = [sp for sp in folie.shapes if sp.has_text_frame and sp.text_frame.text.startswith("Rang ")
              and titel in sp.text_frame.text]
    etikett = [sp for sp in folie.shapes if sp.has_text_frame and sp.text_frame.text.startswith("PRIO ")]
    assert len(rahmen) == 1 and len(etikett) == 1
    return rahmen[0], etikett[0]


def _passt(rahmen) -> bool:
    """Steht der Text bei seiner Schriftgröße vollständig im Rahmen?

    Geschätzt wie in ``zeichnen.passende_groesse``, mit der Zeichenbreite für
    **fette** Schrift (0,58 em) — der Titel ist fett, und genau die zu schmal
    angesetzte Breite ließ ihn in der ersten Fassung unter das Etikett laufen.
    """
    from pptx.util import Emu

    groesse = max(r.font.size.pt for p in rahmen.text_frame.paragraphs for r in p.runs)
    breite_pt = Emu(rahmen.width).pt - 14.4
    hoehe_pt = Emu(rahmen.height).pt - 7.2
    je_zeile = int(breite_pt / (groesse * 0.58))
    zeilen = sum(max(1, -(-len(z) // je_zeile)) for z in rahmen.text_frame.text.split("\n"))
    return zeilen * groesse * 1.2 <= hoehe_pt


@pytest.mark.parametrize("titel", [
    None,  # der längste Titel des Messsatzes aus #167
    "Lebensläufe strukturiert erfassen, Kompetenzprofile aufbauen und mit Projektanforderungen "
    "abgleichen, damit die Vorschlagsliste für jede Ausschreibung automatisch entsteht",
])
def test_der_titel_der_detailfolie_laeuft_nie_unter_das_etikett(messsatz_lauf, titel):
    konzepte, prio, _ = messsatz_lauf
    prio = copy.deepcopy(prio)
    if titel is None:
        eintrag = max(prio["eintraege"], key=lambda e: len(e["titel"]))
    else:
        eintrag = prio["eintraege"][0]
        eintrag["titel"] = titel
    prs = baue_praesentation(konzepte, prio)
    folie = next(f for f in prs.slides if f"potenzial_id: {eintrag['potenzial_id']}" in _notiz(f))

    rahmen, etikett = _titel_und_etikett(folie, eintrag["titel"])
    assert rahmen.left + rahmen.width <= etikett.left, "Titelrahmen reicht unter das Etikett."
    assert eintrag["titel"] in rahmen.text_frame.text, "Der Titel steht nicht vollständig im Rahmen."
    assert rahmen.text_frame.word_wrap is True
    assert _passt(rahmen), "Der Titel passt bei seiner Schriftgröße nicht in den Rahmen."
