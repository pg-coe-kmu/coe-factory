"""
BC2 · Zeichenmittel der Präsentation — Palette und Hilfsfunktionen im KIsult-Schnitt.

**Übernommen aus** ``architektur/archiv/build_template.py`` (#244): die Palette und
``box``/``txt``/``header``/``divider``/``foot``. Was **nicht** übernommen ist: der
Platzhalter-Block ``ph`` (er zeichnete ``{{…}}``-Felder, die ein Mensch füllen
sollte — hier füllt der Code) und das Speichern auf einen absoluten Pfad
außerhalb des Repos.

Die Funktionen nehmen die Folie entgegen statt eine Präsentation aus dem
Modulzustand zu lesen: das Template hielt ``prs`` global, und damit hätte jeder
Aufruf denselben Foliensatz verlängert.
"""

from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

INK = RGBColor(0x1A, 0x1A, 0x2E)
DEEP = RGBColor(0x10, 0x2A, 0x43)
BLUE = RGBColor(0x0F, 0x4C, 0x81)
ACCENT = RGBColor(0x2E, 0x86, 0xDE)
TEAL = RGBColor(0x12, 0x9C, 0x9C)
GREEN = RGBColor(0x1E, 0x88, 0x5E)
AMBER = RGBColor(0xE6, 0x7E, 0x22)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREY = RGBColor(0x5A, 0x6B, 0x7B)
LIGHT = RGBColor(0xEC, 0xF2, 0xF9)
CARD = RGBColor(0xF5, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0x9F, 0xB3, 0xC8)
WARN_FILL = RGBColor(0xFD, 0xED, 0xEC)
HELLBLAU = RGBColor(0xBF, 0xD9, 0xF2)
TITELBLAU = RGBColor(0x8F, 0xB7, 0xE0)

BREITE = Inches(13.333)
HOEHE = Inches(7.5)

__all__ = [
    "ACCENT", "AMBER", "BLUE", "BREITE", "CARD", "DEEP", "GREEN", "GREY", "HELLBLAU", "HOEHE",
    "INK", "LIGHT", "LINE", "RED", "TEAL", "TITELBLAU", "WARN_FILL", "WHITE",
    "box", "divider", "flaeche", "foot", "header", "passende_groesse", "txt",
]


def passende_groesse(text: str, w, h, size: float, mindest: float = 8.0, fett: bool = False) -> float:
    """Die größte Schrift bis ``size``, bei der ``text`` in ``w`` × ``h`` passt — geschätzt.

    **Warum geschätzt und nicht gemessen.** PowerPoint rechnet ``normAutofit``
    erst nach, wenn jemand den Text bearbeitet; eine frisch erzeugte Datei
    zeigt überlaufenden Text also in voller Größe. ``python-pptx`` kann exakt
    messen (``fit_text``), braucht dafür aber eine Schriftdatei auf dem Server.
    Die Schätzung — mittlere Zeichenbreite 0,5 em (fett 0,58 em), Zeilenhöhe 1,2 em — liegt
    eher zu groß und kostet dann eine halbe Stufe Schrift. Abgeschnitten wird
    nie; reicht auch ``mindest`` nicht, bricht der Text weiter um.
    """
    breite_pt = max(w / 12700 - 14.4, 10)
    hoehe_pt = max(h / 12700 - 7.2, 6)
    groesse = size
    while groesse > mindest:
        zeichen_je_zeile = max(int(breite_pt / (groesse * (0.58 if fett else 0.5))), 1)
        zeilen = sum(max(1, -(-len(z) // zeichen_je_zeile)) for z in text.split("\n"))
        if zeilen * groesse * 1.2 <= hoehe_pt:
            return groesse
        groesse -= 0.5
    return mindest


def _absaetze(tf, text: str, size: float, color, bold: bool, align) -> None:
    for i, zeile in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = zeile
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "Calibri"


def box(s, x, y, w, h, text, fill, font=WHITE, size=12, bold=True,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, align=PP_ALIGN.CENTER, line=None, lw=0.75,
        anchor=MSO_ANCHOR.MIDDLE, passend=True):
    """Eine gefüllte Form mit Text. Erste Zeile fett, die übrigen normal und kleiner.

    Langer Text wird **umgebrochen und verkleinert, nie abgeschnitten**
    (``TEXT_TO_FIT_SHAPE``) — die Regel „Überlauf wird umgebrochen, nie begrenzt"
    (#244) gilt auch innerhalb einer Form.
    """
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    if passend:
        size = passende_groesse(text, w, h, size, fett=bold)
    for i, zeile in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = zeile
        r.font.size = Pt(size if i == 0 else max(size - 2, 8))
        r.font.bold = bold if i == 0 else False
        r.font.color.rgb = font
        r.font.name = "Calibri"
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    return sp


def flaeche(s, x, y, w, h, fill, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE):
    """Eine Form ohne Text — Hintergrund, Band, Rahmen."""
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp


def txt(s, x, y, w, h, text, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, schrumpfen=True):
    """Ein Textrahmen ohne Füllung. Schrumpft bei Überlauf, schneidet nie ab."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if schrumpfen:
        size = passende_groesse(text, w, h, size, fett=bold)
    _absaetze(tf, text, size, color, bold, align)
    if schrumpfen:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    return tb


def header(s, kicker, title, color=BLUE, titelbreite=Inches(12.2)):
    """Kopfbalken. ``titelbreite`` schmaler, wenn rechts ein Etikett steht — der
    Titel bricht dann um und schrumpft, statt unter das Etikett zu laufen."""
    flaeche(s, 0, 0, BREITE, Inches(1.1), color)
    flaeche(s, 0, Inches(1.1), BREITE, Inches(0.05), ACCENT)
    txt(s, Inches(0.55), Inches(0.16), Inches(12), Inches(0.3), kicker, 12, HELLBLAU, True)
    txt(s, Inches(0.55), Inches(0.44), titelbreite, Inches(0.62), title, 24, WHITE, True)


def foot(s, n: int, fusszeile: str):
    """Fußzeile. Trägt Mandant, Paket und Fassung — wo sie vom Inhalt nicht abfallen."""
    txt(s, Inches(0.55), Inches(7.08), Inches(11.4), Inches(0.3), fusszeile, 9, GREY)
    txt(s, Inches(12.3), Inches(7.08), Inches(0.7), Inches(0.3), str(n), 9, GREY,
        align=PP_ALIGN.RIGHT)


def divider(s, num: str, title: str, sub: str):
    flaeche(s, 0, 0, BREITE, HOEHE, DEEP)
    flaeche(s, Inches(0.9), Inches(3.7), Inches(2.2), Inches(0.09), ACCENT)
    txt(s, Inches(0.9), Inches(2.4), Inches(4), Inches(1.4), num, 90, RGBColor(0x2E, 0x6B, 0xA8), True,
        schrumpfen=False)
    txt(s, Inches(0.9), Inches(4.0), Inches(11), Inches(0.9), title, 38, WHITE, True)
    txt(s, Inches(0.95), Inches(5.0), Inches(11.5), Inches(1.2), sub, 16, RGBColor(0x9F, 0xB9, 0xD6))
