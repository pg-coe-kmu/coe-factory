"""
BC2 · Der Foliensatz — aus Konzepten und Priorisierung gezeichnet (#257, entschieden in #244).

**Eine reine Funktion.** :func:`baue_praesentation` liest nichts, schreibt
nichts und kennt keine Uhr: sie bekommt die Konzepte und die Priorisierung
eines Laufs in Vertragsform und gibt ein ``Presentation``-Objekt im Speicher
zurück. Abgelegt wird es woanders (:mod:`praesentation.ablage`) — dieselbe
Trennung, die #238 für das Rechenmodell prüfbar gemacht hat.

**Sie trägt keine eigene Information** (Glossar, „Präsentation"). Jeder Satz
auf einer Folie steht in einem der beiden Eingänge oder ist dessen Darstellung;
die Ausgangslage kommt seit v3.1 aus ``priorisierung.ausgangslage`` (#254).
Fehlt etwas, steht dort, *dass* es fehlt — erfunden wird es nicht.

**Der KIsult-Schnitt**, aus ``architektur/archiv/build_template.py``
übernommen und auf v3.1 gehoben:

====  =======================================  ========================================
 Nr   Folie                                    wie oft
====  =======================================  ========================================
  1   Titel (Mandant, Paket, Fassung, Datum)   einmal
  2   Agenda                                   einmal
  3   01 · Ausgangslage                        Trenner
  4   Das Unternehmen                          einmal
  5   Zentrale Herausforderungen               je 6 Herausforderungen eine
  6   02 · Automatisierungspotenziale          Trenner
  7   Überblick                                je 6 Potenziale eine („1 von 2")
  8   Priorisierungsmatrix                     je 9 Potenziale eine
  9   Kernprozess-Abschnitt                    je Kernprozess eine
 10   Potenzial im Detail                      je Potenzial eine, ohne Obergrenze
 11   03 · Kostenschätzung                     Trenner
 12   Kosten und Nutzen je Potenzial           je 7 Potenziale eine
 13   Abschluss                                einmal
====  =======================================  ========================================

**Bewusst entfallen** (#244): die Kausalketten-Folie — „Ursache → Engpass →
verschenkter Ertrag" sind kausale Behauptungen, die aus den Daten nicht folgen —
und die Kacheln „Kunden & Volumen", „Markt & Ziele" sowie die Investitionslogik
mit laufenden Kosten: dafür trägt kein Lauf Daten.

**Gezeigt wird das Ergebnis des Gate 1**, nicht der Vorschlag: nur freigegebene
Potenziale, in der am Gate gesetzten Reihenfolge, falls eine gesetzt ist. Ein
herausgenommenes Potenzial kommt gar nicht vor, auch seine Begründung nicht —
``gate1.nicht_freigegeben[]`` bleibt interne Dokumentation in der Priorisierung.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Mapping, Sequence

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from . import formulierung as F
from .zeichnen import (
    ACCENT, AMBER, BLUE, BREITE, DEEP, GREEN, GREY, HOEHE, INK, LIGHT, LINE, RED, TEAL,
    TITELBLAU, WARN_FILL, WHITE, box, divider, flaeche, foot, header, txt,
)

__all__ = ["KeineFreigabe", "baue_praesentation", "KACHELN_JE_FOLIE", "ZEILEN_MATRIX",
           "ZEILEN_KOSTEN", "HERAUSFORDERUNGEN_JE_FOLIE"]

#: Wie viel auf eine Folie passt. **Grenzen je Folie, nicht je Foliensatz** —
#: was darüber hinausgeht, kommt auf die nächste Folie. Eine „Top 6" wäre
#: derselbe Fehler wie die 0 für eine fehlende Messung: sie sieht vollständig aus.
KACHELN_JE_FOLIE = 6
ZEILEN_MATRIX = 9
ZEILEN_KOSTEN = 7
HERAUSFORDERUNGEN_JE_FOLIE = 6

_KATEGORIE_ANZEIGE = {
    "Quick Win": "Quick Win",
    "Strategisch": "Strategisch",
    "Optional": "Optional",
    "Zurueckgestellt": "Zurückgestellt",
}

_PRIO_FARBE = {"PRIO 1": GREEN, "PRIO 2": AMBER, "PRIO 3": GREY}

_TEXT_FEHLT = "für diesen Lauf nicht vorhanden (der Text entsteht im Erkennungsschritt)"


class KeineFreigabe(ValueError):
    """Die Präsentation entsteht nur nach der Freigabe am Gate 1 (#244).

    Vorher zeigte sie eine Reihenfolge, die der Mensch im nächsten Schritt
    überschreiben darf; bei Ablehnung entsteht sie gar nicht (ADR-007, 2.3).
    """


# ---------------------------------------------------------------------------
# Den Lauf lesen
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class _Eintrag:
    """Ein freigegebenes Potenzial, wie es auf den Folien erscheint."""

    rang: int
    potenzial: Mapping
    eintrag: Mapping

    @property
    def potenzial_id(self) -> str:
        return self.eintrag["potenzial_id"]

    @property
    def titel(self) -> str:
        return self.eintrag["titel"]

    @property
    def kp_id(self) -> str:
        return self.eintrag["kp_id"]

    @property
    def gruppe(self) -> str:
        return self.eintrag["prioritaetsgruppe"]

    @property
    def kategorie(self) -> str:
        return _KATEGORIE_ANZEIGE.get(self.eintrag["kategorie"], self.eintrag["kategorie"])


@dataclass(frozen=True)
class _Lauf:
    priorisierung: Mapping
    konzept_je_kp: dict[str, Mapping]
    eintraege: list[_Eintrag]
    kp_folge: list[str]
    umsortiert: bool

    @property
    def ausgangslage(self) -> Mapping | None:
        return self.priorisierung.get("ausgangslage")

    @property
    def mandant(self) -> str:
        aus = self.ausgangslage
        if aus and aus["unternehmen"].get("name"):
            return aus["unternehmen"]["name"]
        for k in self.konzept_je_kp.values():
            name = (k.get("kontext") or {}).get("unternehmen")
            if name:
                return name
        return f"Mandant {self.priorisierung['company_id']}"

    def je_kp(self, kp_id: str) -> list[_Eintrag]:
        return [e for e in self.eintraege if e.kp_id == kp_id]


def _lies_lauf(konzepte: Sequence[Mapping], priorisierung: Mapping) -> _Lauf:
    gate1 = priorisierung.get("gate1") or {}
    status = gate1.get("status")
    if status != "approved":
        raise KeineFreigabe(
            f"Gate 1 steht auf {status!r}. Die Präsentation entsteht erst nach der Freigabe — "
            "vorher zeigte sie eine Reihenfolge, die der Mensch noch überschreiben darf; "
            "bei Ablehnung entsteht sie gar nicht (#244, ADR-007 2.3)."
        )

    potenzial_je_id = {
        p["potenzial_id"]: p for k in konzepte for p in k.get("potenziale", [])
    }
    konzept_je_kp = {k["kontext"]["kp_id"]: k for k in konzepte if k.get("kontext")}

    eintrag_je_id = {e["potenzial_id"]: e for e in priorisierung["eintraege"]}
    herausgenommen = {n["potenzial_id"] for n in gate1.get("nicht_freigegeben", [])}
    freigegeben = set(gate1.get("approved_potenzial_ids") or eintrag_je_id) - herausgenommen

    folge = list(gate1.get("finale_reihenfolge_potenzial_ids") or [])
    gerechnet = [e["potenzial_id"] for e in sorted(priorisierung["eintraege"],
                                                   key=lambda e: e["potenzialrang"])]
    if not folge:
        folge = gerechnet
    folge = [pid for pid in folge if pid in freigegeben]

    eintraege = [
        _Eintrag(rang=i, potenzial=potenzial_je_id.get(pid, {}), eintrag=eintrag_je_id[pid])
        for i, pid in enumerate(folge, start=1)
    ]

    kp_folge = list(gate1.get("finale_prozessreihenfolge_kp_ids") or [])
    if not kp_folge:
        kp_folge = [r["kp_id"] for r in sorted(priorisierung["prozess_raenge"],
                                               key=lambda r: r["prozessrang"])]
    mit_potenzial = {e.kp_id for e in eintraege}
    kp_folge = [kp for kp in kp_folge if kp in mit_potenzial]

    umsortiert = bool(gate1.get("finale_reihenfolge_potenzial_ids")) and (
        list(gate1["finale_reihenfolge_potenzial_ids"]) != gerechnet
    )
    return _Lauf(priorisierung, konzept_je_kp, eintraege, kp_folge, umsortiert)


def _seiten(liste: list, je_seite: int) -> list[list]:
    """Teilt auf Folien auf. Leer ergibt **eine** leere Seite, nicht keine."""
    if not liste:
        return [[]]
    return [liste[i:i + je_seite] for i in range(0, len(liste), je_seite)]


def _von(i: int, n: int) -> str:
    return f" ({i} von {n})" if n > 1 else ""


# ---------------------------------------------------------------------------
# Der Foliensatz
# ---------------------------------------------------------------------------


class _Satz:
    """Hält die Präsentation und zählt die Seiten — ersetzt das globale ``N`` des Templates."""

    def __init__(self, fusszeile: str) -> None:
        self.prs = Presentation()
        self.prs.slide_width = BREITE
        self.prs.slide_height = HOEHE
        self._leer = self.prs.slide_layouts[6]
        self._fusszeile = fusszeile
        self._n = 0

    def folie(self, *, notiz: str | None = None, fuss: bool = True):
        s = self.prs.slides.add_slide(self._leer)
        self._n += 1
        if fuss:
            foot(s, self._n, self._fusszeile)
        if notiz:
            s.notes_slide.notes_text_frame.text = notiz
        return s


def baue_praesentation(
    konzepte: Sequence[Mapping],
    priorisierung: Mapping,
    *,
    sperrgrund: str | None = None,
    hinweise: Sequence[str] = (),
) -> Presentation:
    """Zeichnet den Foliensatz eines freigegebenen Laufs.

    ``konzepte`` und ``priorisierung`` sind die Liefergegenstände in
    Vertragsform (v3.0/v3.1). Felder, die das LLM schreibt (Beschreibung,
    Vision, Lösungsansatz …), dürfen fehlen — die Oberfläche speist heute aus
    einem Messsatz, der sie nicht trägt; die Folie sagt dann an, dass der Text
    fehlt, statt einen Platzhalter zu zeigen.

    ``sperrgrund`` und ``hinweise`` kommen aus dem Laufkopf (ADR-007 · BC2,
    Nachtrag #305). Sie sagen nichts über den Mandanten, sondern etwas über die
    Herkunft des Laufs, und stehen deshalb auf der Titelfolie, wo sie nicht
    übersehen werden — der Sperrgrund rot, die Hinweise zurückhaltend.

    :raises KeineFreigabe: wenn ``gate1.status`` nicht ``approved`` ist.
    """
    lauf = _lies_lauf(konzepte, priorisierung)
    p = priorisierung
    fusszeile = (
        f"BC2 — Strategic Advisor  |  {lauf.mandant}  ·  Paket {p['paket_id']}  ·  "
        f"Fassung {p['fassung']}"
    )
    satz = _Satz(fusszeile)

    _titel(satz, lauf, sperrgrund, tuple(hinweise))
    _agenda(satz)
    divider(satz.folie(fuss=False), "01", "Ausgangslage",
            "Das Unternehmen und seine zentralen Herausforderungen")
    _unternehmen(satz, lauf)
    _herausforderungen(satz, lauf)
    divider(satz.folie(fuss=False), "02", "Automatisierungspotenziale",
            f"{len(lauf.eintraege)} freigegebene Potenziale in {len(lauf.kp_folge)} Kernprozessen")
    _ueberblick(satz, lauf)
    _matrix(satz, lauf)
    for i, kp_id in enumerate(lauf.kp_folge, start=1):
        _kernprozess(satz, lauf, kp_id, i)
        for e in lauf.je_kp(kp_id):
            _detail(satz, lauf, e)
    _kosten(satz, lauf)
    _abschluss(satz, lauf)
    return satz.prs


# --- 1 Titel ---------------------------------------------------------------


def _titel(
    satz: _Satz, lauf: _Lauf, sperrgrund: str | None, hinweise: tuple[str, ...]
) -> None:
    p = lauf.priorisierung
    g = p["gate1"]
    s = satz.folie(fuss=False)
    flaeche(s, 0, 0, BREITE, HOEHE, DEEP)
    flaeche(s, 0, Inches(4.5), BREITE, Inches(0.08), ACCENT)
    txt(s, Inches(0.9), Inches(1.0), Inches(11.5), Inches(0.5), "KI- & AUTOMATISIERUNGS-WORKSHOP",
        15, TITELBLAU, True)
    txt(s, Inches(0.9), Inches(1.6), Inches(11.5), Inches(1.0), "Zusammenfassung & Potenzialanalyse",
        40, WHITE, True)
    txt(s, Inches(0.9), Inches(2.75), Inches(11.5), Inches(0.6),
        "Ausgangslage · Automatisierungspotenziale · Kostenschätzung", 20,
        TITELBLAU, True)
    freigabe = f"Gate 1 freigegeben am {F.datum(g.get('entschieden_am'))}"
    if g.get("entscheider"):
        freigabe += f" durch {g['entscheider']}"
    zeilen = [
        f"Mandant:      {lauf.mandant}",
        f"Kernprozesse: {', '.join(lauf.kp_folge)}  ·  {len(lauf.eintraege)} freigegebene Potenziale",
        f"Paket {p['paket_id']}  ·  Fassung {p['fassung']}  ·  übergeben am {F.datum(p['uebergeben_am'])}",
        f"{freigabe}  ·  Erstellt von BC2 — Strategic Advisor (Autonomous CoE Factory)",
    ]
    txt(s, Inches(0.9), Inches(4.8), Inches(11.5), Inches(1.5), "\n".join(zeilen), 14, WHITE)
    # Eine Fläche, nicht zwei: zwischen Untertitel und Akzentlinie ist Platz
    # für genau eine. Gibt es einen Sperrgrund, steht er vorn und färbt sie rot.
    if sperrgrund:
        text = " ".join([f"⚠ {sperrgrund}", *hinweise])
        box(s, Inches(0.9), Inches(3.45), Inches(11.5), Inches(0.9), text, WARN_FILL, RED,
            12, True, align=PP_ALIGN.LEFT, line=RED)
    elif hinweise:
        box(s, Inches(0.9), Inches(3.45), Inches(11.5), Inches(0.9), "Hinweis: " + " ".join(hinweise),
            LIGHT, INK, 11, False, align=PP_ALIGN.LEFT, line=LINE)


# --- 2 Agenda --------------------------------------------------------------


def _agenda(satz: _Satz) -> None:
    s = satz.folie()
    header(s, "AGENDA", "Inhalt")
    punkte = [
        ("01", "Ausgangslage", "Das Unternehmen, seine Kernprozesse und zentralen Herausforderungen", BLUE),
        ("02", "Automatisierungspotenziale",
         "Freigegebene Potenziale, Bewertung und Priorisierung (Umsetzungskomplexität × Impact)", GREEN),
        ("03", "Kostenschätzung", "Einsparung, Investitions-Richtwert und Amortisation je Potenzial", AMBER),
    ]
    y = Inches(1.6)
    for n, t, d, c in punkte:
        box(s, Inches(0.6), y, Inches(1.2), Inches(1.2), n, c, size=30, passend=False)
        box(s, Inches(1.95), y, Inches(10.8), Inches(1.2), "", WHITE, INK, 11, False, line=LINE)
        txt(s, Inches(2.2), y + Inches(0.16), Inches(10.3), Inches(0.4), t, 17, c, True)
        txt(s, Inches(2.2), y + Inches(0.62), Inches(10.3), Inches(0.5), d, 12.5, INK)
        y += Inches(1.55)


# --- 4 Unternehmen ---------------------------------------------------------


def _unternehmen(satz: _Satz, lauf: _Lauf) -> None:
    s = satz.folie()
    header(s, "01 · AUSGANGSLAGE", "Das Unternehmen — Daten & Fakten")
    aus = lauf.ausgangslage
    if aus:
        u = aus["unternehmen"]
        mitarbeitende = u["mitarbeitende"]
        unternehmen = "\n".join([
            f"{F.luecke(u['name'])}",
            f"Branche: {F.luecke(u['branche'])}",
            f"Mitarbeitende: {F.luecke(mitarbeitende)}",
            f"Region: {F.luecke(u['region'])}",
        ])
        modell = F.luecke(u["geschaeftsmodell"])
    else:
        grund = ("Liegt in dieser Priorisierung nicht vor: der Mandantensatz reist erst ab "
                 "Vertrag v3.1 mit (#254).")
        unternehmen = modell = grund

    kps = []
    for kp in lauf.kp_folge:
        kurz = ((lauf.konzept_je_kp.get(kp) or {}).get("kontext") or {}).get("prozess_kurzbeschreibung")
        kps.append(f"{kp}: {kurz}" if kurz else kp)
    systeme = sorted({
        sys_
        for kp in lauf.kp_folge
        for sys_ in ((lauf.konzept_je_kp.get(kp) or {}).get("kontext") or {}).get(
            "betroffene_systeme_landschaft", [])
    })
    kacheln = [
        ("UNTERNEHMEN", unternehmen, BLUE),
        ("GESCHÄFTSMODELL", modell, TEAL),
        ("KERNPROZESSE IM PAKET", "\n".join(kps), GREEN),
        ("SYSTEMLANDSCHAFT", ", ".join(systeme) if systeme else
         "Die Konzepte dieses Laufs nennen keine übergreifende Systemlandschaft.", AMBER),
    ]
    xs = [Inches(0.55), Inches(6.75)]
    ys = [Inches(1.5), Inches(4.1)]
    for i, (t, d, c) in enumerate(kacheln):
        x, y = xs[i % 2], ys[i // 2]
        box(s, x, y, Inches(6.0), Inches(2.4), "", WHITE, INK, 11, False, line=c, lw=1.5)
        box(s, x, y, Inches(6.0), Inches(0.55), t, c, size=13, passend=False)
        txt(s, x + Inches(0.2), y + Inches(0.7), Inches(5.6), Inches(1.6), d, 13, INK)


# --- 5 Herausforderungen -----------------------------------------------------


def _herausforderungen(satz: _Satz, lauf: _Lauf) -> None:
    aus = lauf.ausgangslage
    liste = list(aus["herausforderungen"]) if aus else []
    seiten = _seiten(liste, HERAUSFORDERUNGEN_JE_FOLIE)
    kernaussage = (aus or {}).get("kernaussage")
    for n, seite in enumerate(seiten, start=1):
        s = satz.folie()
        header(s, "01 · AUSGANGSLAGE", "Zentrale Herausforderungen" + _von(n, len(seiten)))
        if not aus:
            txt(s, Inches(0.55), Inches(1.6), Inches(12.2), Inches(1.0),
                "Die Herausforderungen liegen in dieser Priorisierung nicht vor: die Ausgangslage "
                "reist erst ab Vertrag v3.1 mit (#254). Sie werden hier nicht nachträglich "
                "zusammengestellt — die Präsentation trägt keine eigene Information.", 14, INK)
            continue
        hoehe_liste = Inches(3.3) if (kernaussage and n == 1) else Inches(5.2)
        zeile_h = int(hoehe_liste / HERAUSFORDERUNGEN_JE_FOLIE)
        for i, h in enumerate(seite):
            y = Inches(1.4) + zeile_h * i
            haeufig = f" ({h['haeufigkeit']})" if h.get("haeufigkeit") else ""
            text = f"{h['beschreibung']}{haeufig}\nAuswirkung: {h['auswirkung']}  ·  {', '.join(h['kp_ids'])}"
            box(s, Inches(0.55), y, Inches(12.25), zeile_h - Inches(0.08), text, LIGHT, INK, 12.5, True,
                align=PP_ALIGN.LEFT, line=LINE)
        if kernaussage and n == 1:
            box(s, Inches(0.55), Inches(4.9), Inches(12.25), Inches(1.5),
                f"KERNAUSSAGE\n{kernaussage}", LIGHT, INK, 13, True, align=PP_ALIGN.LEFT, line=ACCENT)


# --- 7 Überblick -----------------------------------------------------------


def _ueberblick(satz: _Satz, lauf: _Lauf) -> None:
    seiten = _seiten(lauf.eintraege, KACHELN_JE_FOLIE)
    for n, seite in enumerate(seiten, start=1):
        s = satz.folie()
        header(s, "02 · POTENZIALE", "Überblick: freigegebene Potenziale" + _von(n, len(seiten)),
               color=GREEN)
        if lauf.umsortiert:
            begruendung = lauf.priorisierung["gate1"].get("abweichungsbegruendung", "")
            hinweis = f"Reihenfolge am Gate 1 gesetzt, abweichend vom gerechneten Rang. Begründung: {begruendung}"
        else:
            hinweis = "Reihenfolge nach gerechnetem Rang (ADR-006 · BC2), am Gate 1 bestätigt."
        txt(s, Inches(0.55), Inches(1.3), Inches(12.2), Inches(0.45), hinweis, 12, GREY)
        for i, e in enumerate(seite):
            col, row = i % 3, i // 3
            x = Inches(0.55) + Inches(4.18) * col
            y = Inches(1.85) + Inches(2.5) * row
            farbe = _PRIO_FARBE.get(e.gruppe, GREEN)
            box(s, x, y, Inches(4.0), Inches(2.3), "", WHITE, INK, 11, False, line=farbe, lw=1.25)
            box(s, x + Inches(0.18), y + Inches(0.18), Inches(0.6), Inches(0.6), str(e.rang), farbe,
                size=16, shape=MSO_SHAPE.OVAL, passend=False)
            txt(s, x + Inches(0.9), y + Inches(0.15), Inches(3.0), Inches(1.15), e.titel, 13, INK, True)
            txt(s, x + Inches(0.2), y + Inches(1.35), Inches(3.6), Inches(0.85),
                f"{e.kp_id}  ·  {e.gruppe}  ·  {e.kategorie}\n"
                f"Teilprozesse: {', '.join(e.eintrag['betroffene_teilprozess_ids'])}", 11, GREY)


# --- 8 Matrix --------------------------------------------------------------


def _matrix(satz: _Satz, lauf: _Lauf) -> None:
    seiten = _seiten(lauf.eintraege, ZEILEN_MATRIX)
    quadranten = {k: [] for k in ("Quick Win", "Strategisch", "Optional", "Zurückgestellt")}
    for e in lauf.eintraege:
        quadranten.setdefault(e.kategorie, []).append(str(e.rang))
    for n, seite in enumerate(seiten, start=1):
        s = satz.folie()
        header(s, "02 · PRIORISIERUNG", "Priorisierungsmatrix: Umsetzungskomplexität × Impact"
               + _von(n, len(seiten)), color=GREEN)
        txt(s, Inches(0.55), Inches(1.3), Inches(6.6), Inches(0.3),
            "Bewertung je Potenzial, in Rangfolge (Skala 1–10)", 13, BLUE, True)
        koepfe = ["Rang", "Potenzial", "PRIO", "Kompl.", "Impact"]
        breiten = [Inches(0.7), Inches(3.3), Inches(0.9), Inches(0.85), Inches(0.85)]
        xs = [Inches(0.55)]
        for b in breiten[:-1]:
            xs.append(xs[-1] + b)
        for t, x, w in zip(koepfe, xs, breiten):
            box(s, x, Inches(1.7), w, Inches(0.42), t, GREEN, size=10.5, passend=False)
        for i, e in enumerate(seite):
            y = Inches(2.16) + Inches(0.5) * i
            fill = WHITE if i % 2 == 0 else LIGHT
            kompl = e.eintrag["umsetzungskomplexitaet"]
            herkunft = e.eintrag.get("komplexitaet_herkunft")
            zellen = [
                (str(e.rang), PP_ALIGN.CENTER, True),
                (e.titel, PP_ALIGN.LEFT, False),
                (e.gruppe, PP_ALIGN.CENTER, False),
                (f"{kompl}" + (f"\n{herkunft}" if herkunft else ""), PP_ALIGN.CENTER, False),
                (str(e.eintrag["impact"]), PP_ALIGN.CENTER, False),
            ]
            for (t, ausr, fett), x, w in zip(zellen, xs, breiten):
                box(s, x, y, w, Inches(0.48), t, fill, INK, 10, fett, line=LINE, align=ausr)
        _quadranten(s, quadranten)


def _quadranten(s, quadranten: dict[str, list[str]]) -> None:
    mx, my, mw, mh = Inches(7.45), Inches(1.7), Inches(5.3), Inches(4.75)
    box(s, mx, my, mw, mh, "", WHITE, INK, 10, False, line=LINE)
    txt(s, mx, my + Inches(0.05), mw, Inches(0.3), "  hoher Impact ↑   (Ziffern = Rang)", 10, GREY, True)
    txt(s, mx, my + mh - Inches(0.32), mw, Inches(0.3),
        "  Umsetzungskomplexität  (gering → hoch) →", 10, GREY)

    def inhalt(k):
        return ", ".join(quadranten.get(k, [])) or "keines"

    box(s, mx + Inches(0.3), my + Inches(0.45), Inches(2.25), Inches(1.85),
        f"★ ZUERST\nQuick Wins\nRang {inhalt('Quick Win')}", GREEN, size=12)
    box(s, mx + Inches(2.75), my + Inches(0.45), Inches(2.25), Inches(1.85),
        f"Strategisch\nRang {inhalt('Strategisch')}", BLUE, size=12)
    box(s, mx + Inches(0.3), my + Inches(2.45), Inches(2.25), Inches(1.75),
        f"Optional\nRang {inhalt('Optional')}", GREY, size=12)
    box(s, mx + Inches(2.75), my + Inches(2.45), Inches(2.25), Inches(1.75),
        f"Zurückgestellt\nRang {inhalt('Zurückgestellt')}", AMBER, size=12)


# --- 9 Kernprozess ---------------------------------------------------------


def _kernprozess(satz: _Satz, lauf: _Lauf, kp_id: str, nr: int) -> None:
    s = satz.folie(notiz=f"kp_id: {kp_id}")
    header(s, f"02 · KERNPROZESS {nr} VON {len(lauf.kp_folge)}", f"Kernprozess {kp_id}", color=GREEN)
    kontext = (lauf.konzept_je_kp.get(kp_id) or {}).get("kontext") or {}
    kurz = kontext.get("prozess_kurzbeschreibung") or f"Kurzbeschreibung: {_TEXT_FEHLT}."
    txt(s, Inches(0.55), Inches(1.35), Inches(12.2), Inches(0.9), kurz, 15, INK)

    txt(s, Inches(0.55), Inches(2.4), Inches(6), Inches(0.35), "SCHMERZPUNKTE", 12, BLUE, True)
    schmerz = kontext.get("hauptschmerzpunkte") or []
    zeilen = [f"• {sp['beschreibung']} — {sp['auswirkung']}" for sp in schmerz] or [
        f"Schmerzpunkte: {_TEXT_FEHLT}."]
    txt(s, Inches(0.55), Inches(2.8), Inches(6.1), Inches(3.9), "\n".join(zeilen), 13, INK)

    txt(s, Inches(6.95), Inches(2.4), Inches(6), Inches(0.35), "FREIGEGEBENE POTENZIALE", 12, GREEN, True)
    pots = [f"Rang {e.rang}  ·  {e.gruppe}  ·  {e.titel}" for e in lauf.je_kp(kp_id)]
    txt(s, Inches(6.95), Inches(2.8), Inches(5.85), Inches(3.9), "\n".join(pots), 13, INK)


# --- 10 Detail -------------------------------------------------------------


def _text(pot: Mapping, feld: str, bezeichnung: str) -> str:
    wert = pot.get(feld)
    if isinstance(wert, str) and wert.strip():
        return wert.strip()
    return f"{bezeichnung}: {_TEXT_FEHLT}."


def _detail(satz: _Satz, lauf: _Lauf, e: _Eintrag) -> None:
    pot = e.potenzial
    s = satz.folie(notiz=f"potenzial_id: {e.potenzial_id}\nkp_id: {e.kp_id}")
    # Das Etikett steht rechts im Kopfbalken; der Titel endet 0,15 Zoll davor und bricht um.
    etikett_x = Inches(10.2)
    header(s, f"02 · {e.kp_id} · POTENZIAL", f"Rang {e.rang}   {e.titel}", color=GREEN,
           titelbreite=etikett_x - Inches(0.55) - Inches(0.15))
    box(s, etikett_x, Inches(0.25), Inches(2.85), Inches(0.55), f"{e.gruppe} · {e.kategorie}",
        _PRIO_FARBE.get(e.gruppe, AMBER), size=12)

    links_x, links_w = Inches(0.55), Inches(6.1)
    txt(s, links_x, Inches(1.3), links_w, Inches(0.3), "HEUTIGER PROZESS", 12, BLUE, True)
    systeme = ", ".join(sys_.get("name", "") for sys_ in pot.get("betroffene_systeme", [])) or F.LUECKE
    heute = "\n".join([
        _text(pot, "beschreibung", "Beschreibung"),
        F.manueller_aufwand(pot),
        f"Teilprozesse: {', '.join(e.eintrag['betroffene_teilprozess_ids'])}  ·  Systeme: {systeme}",
    ])
    txt(s, links_x, Inches(1.62), links_w, Inches(2.35), heute, 11, INK)

    txt(s, links_x, Inches(4.05), links_w, Inches(0.3), "AUTOMATISIERUNGSPOTENZIAL", 12, GREEN, True)
    ansatz = (pot.get("potenzielle_loesung") or {}).get("ansatz")
    soll = "\n".join([
        _text(pot, "to_be_vision", "Soll-Vision"),
        f"Lösungsansatz: {ansatz}" if ansatz else f"Lösungsansatz: {_TEXT_FEHLT}.",
        F.automatisierungsgrad(pot),
    ])
    txt(s, links_x, Inches(4.37), links_w, Inches(2.6), soll, 11, INK)

    rechts_x, rechts_w = Inches(6.95), Inches(5.85)
    impact_mon = e.eintrag.get("impact_monetaer")
    monetaer = f"monetär {impact_mon}" if impact_mon is not None else "monetär: keine Wertaussage"
    kompl_herkunft = e.eintrag.get("komplexitaet_herkunft") or F.LUECKE
    bewertungen = [
        ("EINSPARUNG", "\n".join([F.einsparung(pot), F.investition(pot), F.amortisation(pot)]), GREEN),
        ("IMPACT", f"{e.eintrag['impact']} von 10  —  {monetaer}, Nutzwert "
                   f"{F.zahl(e.eintrag['nutzwert_mittel'], 1)}", BLUE),
        ("UMSETZUNGSKOMPLEXITÄT", f"{e.eintrag['umsetzungskomplexitaet']} von 10  —  {kompl_herkunft}", AMBER),
    ]
    hoehen = [Inches(1.25), Inches(0.62), Inches(0.62)]
    y = Inches(1.3)
    for (t, d, c), h in zip(bewertungen, hoehen):
        box(s, rechts_x, y, rechts_w, h, "", WHITE, INK, 10, False, line=c, lw=1.25)
        txt(s, rechts_x + Inches(0.15), y + Inches(0.02), Inches(2.6), Inches(0.3), t, 10.5, c, True)
        txt(s, rechts_x + Inches(0.15), y + Inches(0.28), rechts_w - Inches(0.3), h - Inches(0.3),
            d, 11, INK, True)
        y += h + Inches(0.08)

    q = pot.get("querschnitte") or {}
    quer = []
    quer.append(f"Zukunftssicherheit: {q['zukunftssicherheit']}" if q.get("zukunftssicherheit")
                else f"Zukunftssicherheit: {_TEXT_FEHLT}")
    if q.get("reifegrad"):
        quer.append(f"Reifegrad: {q['reifegrad']}")
    if q.get("umsatzpotenzial"):
        quer.append(f"Umsatzpotenzial: {q['umsatzpotenzial']}")
    if q.get("abhaengigkeiten"):
        quer.append(f"Abhängigkeiten: {'; '.join(q['abhaengigkeiten'])}")
    txt(s, rechts_x, Inches(4.08), rechts_w, Inches(0.3), "QUERSCHNITTE (gehen in keine Rechnung ein)",
        11, GREY, True)
    txt(s, rechts_x, Inches(4.36), rechts_w, Inches(1.35), "\n".join(quer), 10.5, INK)

    voraus = pot.get("voraussetzungen")
    txt(s, rechts_x, Inches(5.75), rechts_w, Inches(0.3), "VORAUSSETZUNGEN", 11, GREY, True)
    vtext = "; ".join(voraus) if voraus else f"Voraussetzungen: {_TEXT_FEHLT}."
    hinweise = pot.get("hinweise") or []
    if hinweise:
        vtext += "\nHinweise: " + "; ".join(f"{h['art']}: {h['text']}" for h in hinweise)
    txt(s, rechts_x, Inches(6.03), rechts_w, Inches(0.98), vtext, 10.5, INK)


# --- 12 Kosten -------------------------------------------------------------


def _kosten(satz: _Satz, lauf: _Lauf) -> None:
    mit_wert = [e for e in lauf.eintraege if F.hat_wertaussage(e.potenzial)]
    keiner = not mit_wert
    if keiner:
        unter = (f"Für keines der {len(lauf.eintraege)} Potenziale liegt eine Wertaussage vor — "
                 "statt leerer Tabellen steht je Potenzial der Grund.")
    else:
        unter = "Spannen statt Punktwerte, die Herkunft im selben Satz"
    divider(satz.folie(fuss=False), "03", "Kostenschätzung", unter)

    if keiner:
        koepfe = ["Rang", "Potenzial", "Warum keine Wertaussage", "Aufwand"]
        breiten = [Inches(0.8), Inches(3.6), Inches(5.6), Inches(2.25)]
    else:
        koepfe = ["Rang", "Potenzial", "Einsparung je Jahr", "Investition", "Amortisation"]
        breiten = [Inches(0.8), Inches(3.1), Inches(4.0), Inches(2.6), Inches(1.75)]
    xs = [Inches(0.55)]
    for b in breiten[:-1]:
        xs.append(xs[-1] + b)

    seiten = _seiten(lauf.eintraege, ZEILEN_KOSTEN)
    for n, seite in enumerate(seiten, start=1):
        s = satz.folie()
        header(s, "03 · KOSTENSCHÄTZUNG", "Kosten und Nutzen je Potenzial" + _von(n, len(seiten)),
               color=AMBER)
        for t, x, w in zip(koepfe, xs, breiten):
            box(s, x, Inches(1.45), w, Inches(0.5), t, AMBER, size=11, passend=False)
        for i, e in enumerate(seite):
            y = Inches(1.98) + Inches(0.66) * i
            fill = WHITE if i % 2 == 0 else LIGHT
            pot = e.potenzial
            if keiner:
                grund = F.einsparung(pot).removeprefix("Keine Wertaussage: ")
                zellen = [str(e.rang), e.titel, grund,
                          F.investition(pot).removeprefix("Investition: ")]
            else:
                zellen = [
                    str(e.rang), e.titel, F.einsparung(pot),
                    F.investition(pot).removeprefix("Investition: "),
                    F.amortisation(pot).removeprefix("Amortisation "),
                ]
            for j, (t, x, w) in enumerate(zip(zellen, xs, breiten)):
                box(s, x, y, w, Inches(0.62), t, fill, INK, 10.5, j == 0, line=LINE,
                    align=PP_ALIGN.CENTER if j == 0 else PP_ALIGN.LEFT)
        txt(s, Inches(0.55), Inches(6.62), Inches(12.2), Inches(0.42),
            "Einsparung: Eckenrechnung aus der Bandbreite der Dauer und dem Korridor des "
            "Automatisierungsgrads (ADR-006 · BC2). Investition: Richtwert aus dem geschätzten Aufwand "
            "in Personentagen — konkrete Festpreise erst nach Detaillierung.", 9.5, GREY)


# --- 13 Abschluss ----------------------------------------------------------


def _abschluss(satz: _Satz, lauf: _Lauf) -> None:
    p = lauf.priorisierung
    s = satz.folie(fuss=False)
    flaeche(s, 0, 0, BREITE, HOEHE, DEEP)
    flaeche(s, Inches(0.9), Inches(3.5), Inches(2.2), Inches(0.09), ACCENT)
    txt(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.0), "Vielen Dank", 44, WHITE, True)
    kernaussage = (lauf.ausgangslage or {}).get("kernaussage")
    if kernaussage:
        txt(s, Inches(0.9), Inches(3.8), Inches(11.5), Inches(1.6), kernaussage, 16, WHITE)
    grundlage = (
        f"Grundlage: Priorisierung und {len(lauf.konzept_je_kp)} Konzept(e) des Analyselaufs "
        f"Paket {p['paket_id']}, Fassung {p['fassung']}, übergeben am {F.datum(p['uebergeben_am'])} "
        f"— Vertrag BC2 → BC3 v{p['schema_version']}. Diese Präsentation trägt keine Information, "
        "die nicht dort steht."
    )
    txt(s, Inches(0.9), Inches(5.7), Inches(11.5), Inches(0.7), grundlage, 11, TITELBLAU)
    txt(s, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.4),
        "BC2 — Strategic Advisor · Autonomous CoE Factory", 12, TITELBLAU)
